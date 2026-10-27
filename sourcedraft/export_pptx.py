"""Fill a .pptx template: student points on slides, instructor notes in speaker notes, source line per slide."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def _sources(claims):
    return sorted({s for c in claims for s in c["sources"]})


def _textbox(slide, left, top, width, height, text, size, color, bold=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.text_frame.word_wrap = True
    run = box.text_frame.paragraphs[0].add_run()
    run.text, run.font.size, run.font.bold = text, Pt(size), bold
    run.font.color.rgb = RGBColor(*color)
    return box


def _watermark(prs, slide):
    box = _textbox(slide, 0, int(prs.slide_height * 0.35), prs.slide_width, Inches(2), "DRAFT", 96, (255, 190, 190), True)
    box.text_frame.paragraphs[0].alignment = 2  # centre
    box.rotation = -25


def export(data, out, template=None, final=False):
    prs = Presentation(template) if template else Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = data["title"]
    slides = [title]
    for s in data["slides"]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = s["title"]
        slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
        body = slide.placeholders[1]
        left, top, width = body.left, body.top, body.width
        body.left, body.top, body.width = left, top, width
        body.height = prs.slide_height - top - Inches(0.8)  # keep clear of the source line
        frame = body.text_frame
        frame.text = s["points"][0]["text"] if s["points"] else ""
        for p in s["points"][1:]:
            frame.add_paragraph().text = p["text"]
        for para in frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(22)
        notes = s.get("notes", [])
        slide.notes_slide.notes_text_frame.text = "\n".join(n["text"] for n in notes)
        line = "Sources: " + ", ".join(_sources(s["points"] + notes))
        _textbox(slide, Inches(0.4), prs.slide_height - Inches(0.6), prs.slide_width - Inches(0.8), Inches(0.4), line, 9, (110, 110, 110))
        slides.append(slide)
    if not final:
        for slide in slides:
            _watermark(prs, slide)
    prs.save(out)
