"""Fill a .docx template: student vignette and a separate teaching note, both from the same facts."""
from docx import Document


def _doc(template, final):
    doc = Document(template) if template else Document()
    if not final:
        doc.sections[0].header.paragraphs[0].text = "DRAFT - not approved"
    return doc


def _para(doc, claim):
    doc.add_paragraph(claim["text"])
    doc.add_paragraph("Sources: " + ", ".join(claim["sources"])).runs[0].italic = True


def export(data, out_dir, template=None, final=False):
    """Writes vignette.docx and teaching-note.docx into out_dir (prefixed draft-/final-)."""
    prefix = "final-" if final else "draft-"
    v = _doc(template, final)
    v.add_heading(data["title"], 0)
    for c in data["vignette"]:
        v.add_paragraph(c["text"])
    v.add_heading("Decision point", 1)
    v.add_paragraph(data["decision_point"]["text"])
    v.add_heading("Questions", 1)
    for q in data["questions"]:
        v.add_paragraph(q, style="List Number")
    v.save(out_dir / f"{prefix}vignette.docx")

    n = _doc(template, final)
    n.add_heading(f"Teaching note: {data['title']}", 0)
    n.add_heading("Facts used", 1)
    for c in data["vignette"] + [data["decision_point"]]:
        _para(n, c)
    n.add_heading("Teaching note", 1)
    for c in data["teaching_note"]:
        _para(n, c)
    n.save(out_dir / f"{prefix}teaching-note.docx")
