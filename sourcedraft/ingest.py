"""Turn source files into chunks: {id, file, text}. Chunk id = <file>#<slide|page|section>."""
import re
from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader


def _pptx(path):
    for n, slide in enumerate(Presentation(path).slides, 1):
        parts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
        if slide.has_notes_slide:
            parts.append("Notes: " + slide.notes_slide.notes_text_frame.text)
        yield f"slide{n}", "\n".join(parts)


def _pdf(path):
    for n, page in enumerate(PdfReader(path).pages, 1):
        yield f"p{n}", page.extract_text() or ""


def _sections(lines):
    """Group (heading, text) pairs from (is_heading, text) lines."""
    head, buf = "intro", []
    for is_head, text in lines:
        if is_head:
            yield head, "\n".join(buf)
            head, buf = text, []
        else:
            buf.append(text)
    yield head, "\n".join(buf)


def _docx(path):
    lines = [(p.style.name.startswith("Heading"), p.text) for p in Document(path).paragraphs]
    return _sections(lines)


def _md(path):
    lines = [(bool(re.match(r"#+\s", l)), l.lstrip("# ").strip() if l.startswith("#") else l)
             for l in Path(path).read_text().splitlines()]
    return _sections(lines)


READERS = {".pptx": _pptx, ".pdf": _pdf, ".docx": _docx, ".md": _md, ".txt": _md}


def ingest_dir(src):
    chunks = []
    for path in sorted(Path(src).rglob("*")):
        reader = READERS.get(path.suffix.lower())
        if not reader:
            continue
        seen = {}
        for label, text in reader(path):
            if not text.strip():
                continue
            seen[label] = seen.get(label, 0) + 1
            suffix = f"-{seen[label]}" if seen[label] > 1 else ""
            chunks.append({"id": f"{path.name}#{label}{suffix}", "file": path.name, "text": text.strip()})
    return chunks
