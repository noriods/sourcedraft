import json
import shutil
import sys
from pathlib import Path

import pytest
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from sourcedraft import ROOT, provider, runs
from sourcedraft.validate import validate

DECK = {"title": "Planning", "slides": [{
    "title": "What is planning", "points": [{"text": "Planning sets goals.", "sources": ["a.md#Intro"]}],
    "notes": [{"text": "Ask for an example.", "sources": ["a.md#Intro"]}]}]}


@pytest.fixture
def project(tmp_path, monkeypatch):
    monkeypatch.setattr(runs, "ROOT", tmp_path)
    shutil.copytree(ROOT / "recipes", tmp_path / "recipes")
    (tmp_path / "templates").mkdir()
    for f in (ROOT / "templates").iterdir():
        shutil.copy(f, tmp_path / "templates")
    p = tmp_path / "projects" / "t"
    (p / "sources").mkdir(parents=True)
    (p / "sources" / "a.md").write_text("# Intro\nPlanning sets goals for the team.\n")
    (p / "project.yaml").write_text("provider: {base_url: 'http://localhost:1', model: fake, api_key_env: X}\n")
    runs.ingest_project("t")
    return tmp_path


def fake(data):
    return lambda cfg, messages, privacy="any": json.dumps(data)


def test_validator_rejects_unknown_and_uncited():
    assert validate(DECK, {"a.md#Intro"})["ok"]
    assert not validate(DECK, {"other#1"})["ok"]
    bad = {"slides": [{"points": [{"text": "x", "sources": []}]}]}
    assert not validate(bad, {"a.md#Intro"})["ok"]


def test_pptx_has_notes_and_source_line(project):
    d = runs.draft("t", "teaching-deck", "planning", chat=fake(DECK))
    prs = Presentation(d / "draft.pptx")
    slide = prs.slides[1]
    assert "Ask for an example." in slide.notes_slide.notes_text_frame.text
    assert any("Sources: a.md#Intro" in s.text_frame.text for s in slide.shapes if s.has_text_frame)


def test_invalid_model_output_fails_run(project):
    bad = {**DECK, "slides": [{**DECK["slides"][0], "points": [{"text": "x", "sources": ["nope#1"]}]}]}
    d = runs.draft("t", "teaching-deck", "planning", chat=fake(bad))
    assert runs.read(d, "status.json")["status"] == "failed"
    assert not (d / "draft.pptx").exists()


def test_final_blocked_without_approval_then_allowed(project):
    d = runs.draft("t", "teaching-deck", "planning", chat=fake(DECK))
    with pytest.raises(PermissionError):
        runs.export(d, final=True)
    assert not (d / "final.pptx").exists()
    with pytest.raises(ValueError):
        runs.approve(d, " ")
    runs.approve(d, "Dr Smith")
    assert (d / "final.pptx").exists()
    assert runs.read(d, "status.json")["status"] == "final"
    assert "approved" in (d / "audit.jsonl").read_text()


def test_vignette_docx(project):
    v = {"title": "T", "vignette": [{"text": "s", "sources": ["a.md#Intro"]}],
         "decision_point": {"text": "d", "sources": ["a.md#Intro"]}, "questions": ["q?"],
         "teaching_note": [{"text": "n", "sources": ["a.md#Intro"]}]}
    d = runs.draft("t", "case-vignette", "planning", chat=fake(v))
    assert (d / "draft-vignette.docx").exists() and (d / "draft-teaching-note.docx").exists()


def test_local_only_blocks_remote():
    with pytest.raises(PermissionError):
        provider.check_privacy("https://openrouter.ai/api/v1", "local-only")
    provider.check_privacy("http://localhost:11434/v1", "local-only")
